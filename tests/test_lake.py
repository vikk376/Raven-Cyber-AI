import io
import json

import pandas as pd
import pytest

from src.services import lake, lake_agent

CSV = b"user,country,score\nalice,CA,5\nbob,US,7\ncarol,CA,9\n"


def test_csv_becomes_queryable_table(db):
    (ds,) = lake.ingest("a@x.com", "Sign ins.csv", CSV, description="d", tags="identity, logs")
    assert ds["kind"] == "table" and ds["rows"] == 3 and ds["table_name"] == "sign_ins"
    assert ds["tag_list"] == ["identity", "logs"]
    df = lake.query("SELECT country, COUNT(*) AS n FROM sign_ins GROUP BY country ORDER BY n DESC")
    assert df.iloc[0].tolist() == ["CA", 2]
    assert len(lake.preview(ds)) == 3
    assert lake.raw_path(ds).read_bytes() == CSV


def test_json_array_is_flattened_and_nested_doc_is_kept(db):
    (t,) = lake.ingest("a@x.com", "events.json", json.dumps(
        [{"id": 1, "user": {"name": "alice"}, "tags": ["a", "b"]}, {"id": 2, "user": {"name": "bob"}, "tags": []}]).encode())
    cols = [c["name"] for c in t["columns"]]
    assert "user_name" in cols or "user.name".replace(".", "_") in cols
    wrapped = json.dumps({"meta": {"v": 1}, "items": [{"k": 1}, {"k": 2}]}).encode()
    (w,) = lake.ingest("a@x.com", "wrapped.json", wrapped)
    assert w["kind"] == "table" and w["rows"] == 2
    (d,) = lake.ingest("a@x.com", "config.json", json.dumps({"a": {"b": [1, 2, 3]}, "name": "x"}).encode())
    assert d["kind"] == "document" and d["rows"] == 0
    assert "name" in lake.document_outline(d)
    assert json.loads(lake.read_document(d))["name"] == "x"


def test_jsonl_and_excel_sheets(db):
    (j,) = lake.ingest("a@x.com", "log.jsonl", b'{"a":1}\n{"a":2}\n')
    assert j["rows"] == 2
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame({"x": [1, 2]}).to_excel(xw, sheet_name="One", index=False)
        pd.DataFrame({"y": [3]}).to_excel(xw, sheet_name="Two", index=False)
    made = lake.ingest("a@x.com", "book.xlsx", buf.getvalue())
    assert [m["rows"] for m in made] == [2, 1] and made[0]["name"] == "book - One"


def test_duplicate_names_get_suffix_and_bad_files_are_rejected(db):
    a = lake.ingest("a@x.com", "d.csv", CSV)[0]
    b = lake.ingest("a@x.com", "d.csv", CSV)[0]
    assert a["name"] != b["name"] and a["table_name"] != b["table_name"]
    for name, data in (("x.exe", b"MZ"), ("empty.csv", b""), ("bad.json", b"{nope"), ("lines.jsonl", b"1\n2\n")):
        with pytest.raises(lake.LakeError):
            lake.ingest("a@x.com", name, data)


def test_columns_are_made_sql_safe(db):
    (ds,) = lake.ingest("a@x.com", "w.csv", b"First Name,1st,First Name\na,b,c\n")
    assert [c["name"] for c in ds["columns"]] == ["First_Name", "c_1st", "First_Name_1"]


@pytest.mark.parametrize("sql", [
    "DROP TABLE t", "DELETE FROM t", "SELECT 1; DROP TABLE t", "PRAGMA table_info(t)", "ATTACH DATABASE 'x' AS y",
    "SELECT * FROM sqlite_master", "UPDATE t SET a=1", "WITH x AS (SELECT 1) DELETE FROM t"])
def test_query_is_read_only(db, sql):
    lake.ingest("a@x.com", "t.csv", b"a\n1\n")
    with pytest.raises(lake.LakeError):
        lake.query(sql)
    assert lake.query("SELECT COUNT(*) AS n FROM t")["n"][0] == 1


def test_allowed_tables_restrict_reads(db):
    lake.ingest("a@x.com", "one.csv", b"a\n1\n")
    lake.ingest("a@x.com", "two.csv", b"b\n2\n")
    assert len(lake.query("SELECT * FROM one", allowed_tables={"one"})) == 1
    with pytest.raises(lake.LakeError):
        lake.query("SELECT * FROM two", allowed_tables={"one"})


def test_lake_cannot_see_platform_database(db):
    lake.ingest("a@x.com", "t.csv", b"a\n1\n")
    with pytest.raises(lake.LakeError):
        lake.query("SELECT * FROM users")


def test_delete_removes_table_and_file(db):
    (ds,) = lake.ingest("a@x.com", "t.csv", b"a\n1\n")
    lake.delete(ds["id"])
    assert lake.get(ds["id"]) is None and not list((lake.root() / "raw").iterdir())
    with pytest.raises(lake.LakeError):
        lake.query("SELECT * FROM t")


def test_stats_and_schema_text(db):
    lake.ingest("a@x.com", "t.csv", CSV)
    lake.ingest("a@x.com", "c.json", b'{"a": {"b": 1}}')
    s = lake.stats()
    assert (s["datasets"], s["tables"], s["documents"], s["rows"]) == (2, 1, 1, 3)
    txt = lake.schema_text()
    assert "table t" in txt and "document c" in txt


def test_agent_tools_are_scoped_and_read_only(db):
    a = lake.ingest("a@x.com", "alpha.csv", CSV)[0]
    lake.ingest("a@x.com", "beta.csv", b"z\n1\n")
    d = lake.ingest("a@x.com", "doc.json", b'{"k": "v"}')[0]
    tools = {t.name: t for t in lake_agent.make_tools([a, d])}
    assert "alpha" in tools["list_datasets"].invoke({}) and "beta" not in tools["list_datasets"].invoke({})
    assert "alice" in tools["run_sql"].invoke({"query": "SELECT * FROM alpha"})
    assert tools["run_sql"].invoke({"query": "SELECT * FROM beta"}).startswith("Error")
    assert tools["run_sql"].invoke({"query": "DROP TABLE alpha"}).startswith("Error")
    assert "country" in tools["describe_table"].invoke({"table": "alpha"})
    assert '"k"' in tools["read_document"].invoke({"name": "doc"})
    assert "Unknown" in tools["read_document"].invoke({"name": "nope"})


def test_ask_with_no_datasets_selected(db):
    answer, tools = lake_agent.ask("hi", [], [])
    assert "no datasets" in answer.lower() and tools == []


# ---- RAG: PDF, Word and text
def _pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(40, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def _docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_paragraph("Incident response policy. Compromised accounts must be disabled within one hour.")
    d.add_paragraph("Password resets require manager approval and a ticket number.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Severity", "SLA"
    t.cell(1, 0).text, t.cell(1, 1).text = "Critical", "15 minutes"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_is_chunked_per_page_and_searchable(db):
    pdf = _pdf(["Ransomware playbook\nIsolate the infected host immediately.", "Phishing playbook\nReset the password of the user and revoke sessions."])
    (ds,) = lake.ingest("a@x.com", "playbooks.pdf", pdf)
    assert ds["kind"] == "text" and ds["info"]["pages"] == 2 and ds["rows"] >= 2
    hits = lake.search_passages("how do I reset a password after phishing", [ds])
    assert hits and hits[0]["page"] == 2 and hits[0]["document"] == "playbooks"
    assert lake.search_passages("isolate infected host", [ds])[0]["page"] == 1
    assert lake.search_passages("zebra", [ds]) == []


def test_word_and_text_documents(db):
    (w,) = lake.ingest("a@x.com", "policy.docx", _docx())
    assert w["kind"] == "text" and w["info"]["format"] == "docx"
    hit = lake.search_passages("critical SLA minutes", [w])[0]
    assert "15 minutes" in hit["text"] and hit["page"] is None
    (t,) = lake.ingest("a@x.com", "notes.md", b"# Notes\n\nThe VPN gateway is vpn01.corp.com.\n")
    assert "vpn01" in lake.search_passages("vpn gateway", [t])[0]["text"]


def test_scanned_or_empty_documents_are_rejected(db):
    with pytest.raises(lake.LakeError):
        lake.ingest("a@x.com", "blank.pdf", _pdf([""]))
    with pytest.raises(lake.LakeError):
        lake.ingest("a@x.com", "broken.pdf", b"not a pdf")
    with pytest.raises(lake.LakeError):
        lake.ingest("a@x.com", "broken.docx", b"not a docx")


def test_search_is_scoped_to_selected_documents_and_delete_cleans_up(db):
    (a,) = lake.ingest("a@x.com", "a.txt", b"The firewall blocks port 445.")
    (b,) = lake.ingest("a@x.com", "b.txt", b"The firewall vendor is Palo Alto.")
    assert [h["document"] for h in lake.search_passages("firewall", [a])] == ["a"]
    assert {h["document"] for h in lake.search_passages("firewall", [a, b])} == {"a", "b"}
    lake.delete(a["id"])
    assert [h["document"] for h in lake.search_passages("firewall")] == ["b"]
    assert lake.stats()["passages"] == 1


def test_long_text_is_split_into_overlapping_passages(db):
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 120 for i in range(12)).encode()
    (ds,) = lake.ingest("a@x.com", "long.txt", text)
    assert ds["rows"] > 4
    from src.services import rag
    rows = rag.chunk_rows(lake._lake_db(), ds["id"])
    assert all(len(r["text"]) <= 1400 for r in rows)


def test_sql_cannot_read_rag_tables(db):
    lake.ingest("a@x.com", "a.txt", b"hello world")
    with pytest.raises(lake.LakeError):
        lake.query("SELECT * FROM rag_chunks")


def test_hybrid_search_with_embeddings(db, monkeypatch):
    import numpy as np

    from src.services import rag
    vocab = ["ransomware", "phishing", "password", "isolate", "mailbox"]

    def vec(t):
        v = np.array([t.lower().count(w) for w in vocab], dtype=float) + 0.01
        return (v / np.linalg.norm(v)).tolist()

    monkeypatch.setattr(rag, "embeddings_enabled", lambda: True)
    monkeypatch.setattr(rag, "embed_texts", lambda texts: [vec(t) for t in texts])
    monkeypatch.setattr(rag, "embed_query", vec)
    (ds,) = lake.ingest("a@x.com", "p.txt", b"Ransomware response: isolate the host.\n\nPhishing response: reset password, purge the mailbox.")
    assert ds["info"]["semantic"] is True
    hit = lake.search_passages("what to do about phishing mailbox", [ds])[0]
    assert "Phishing" in hit["text"] and "semantic" in hit["mode"]
    assert "semantic" in lake.search_mode().lower()


def test_embedding_failure_falls_back_to_keywords(db, monkeypatch):
    from src.services import rag
    monkeypatch.setattr(rag, "embeddings_enabled", lambda: True)
    monkeypatch.setattr(rag, "embed_texts", lambda texts: (_ for _ in ()).throw(RuntimeError("deployment not found")))
    (ds,) = lake.ingest("a@x.com", "p.txt", b"The backup server is bak01.")
    assert ds["info"]["semantic"] is False and "keyword" in ds["info"]["note"].lower()
    assert lake.search_passages("backup server", [ds])


def test_agent_search_documents_tool_returns_citations(db):
    (ds,) = lake.ingest("a@x.com", "pol.pdf", _pdf(["Access policy", "Break glass accounts are reviewed monthly."]))
    tools = {t.name: t for t in lake_agent.make_tools([ds])}
    out = tools["search_documents"].invoke({"query": "break glass review"})
    assert "[pol, p.2]" in out and "monthly" in out
    assert "pol" in tools["list_datasets"].invoke({}) and "search_documents" in tools["list_datasets"].invoke({})
