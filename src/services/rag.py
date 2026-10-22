"""Retrieval for Data Lake documents (PDF, Word, text): extract, chunk, index, search.

Search is hybrid. SQLite FTS5 (BM25 keyword ranking) always works. When an Azure OpenAI embeddings deployment is set
(AZURE_OPENAI_EMBEDDING_DEPLOYMENT) passages are also embedded, and keyword and semantic rankings are merged with
reciprocal rank fusion. Everything is stored in data/lake/lake.db.
"""
import io
import re
import sqlite3
from pathlib import Path

import numpy as np

from src.core import config

TEXT_EXT = (".pdf", ".docx", ".txt", ".md")
CHUNK_CHARS = 1000
OVERLAP_CHARS = 150
MAX_CHUNKS = 20000
EMBED_BATCH = 64


class RagError(Exception):
    pass


# ------------------------------------------------------------------ extraction
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    raise RagError("Cannot decode the file as text.")


def extract(ext: str, data: bytes) -> list[tuple[int | None, str]]:
    """Return [(page_or_None, text)] with empty pages dropped."""
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise RagError("The PDF is password protected.")
            pages = [(i + 1, (pg.extract_text() or "")) for i, pg in enumerate(reader.pages)]
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            pages = [(None, "\n".join(parts))]
        else:
            pages = [(None, _decode(data))]
    except RagError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise RagError(f"Could not read this {ext} file: {str(exc)[:200]}") from exc
    pages = [(n, re.sub(r"[ \t]+", " ", t).strip()) for n, t in pages]
    pages = [(n, t) for n, t in pages if t]
    if not pages:
        raise RagError("No text found. A scanned PDF or an image-only document needs OCR first.")
    return pages


def chunk(pages: list[tuple[int | None, str]]) -> list[dict]:
    """Split into overlapping passages on paragraph and sentence boundaries, never across pages."""
    out: list[dict] = []
    for page, text in pages:
        paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[A-Z0-9•\-])", text) if p.strip()]
        buf = ""
        for para in paras:
            for piece in _split_long(para):
                if buf and len(buf) + len(piece) + 1 > CHUNK_CHARS:
                    out.append({"page": page, "text": buf})
                    buf = buf[-OVERLAP_CHARS:].split(" ", 1)[-1] + " " + piece
                else:
                    buf = (buf + " " + piece).strip()
        if buf:
            out.append({"page": page, "text": buf})
        if len(out) > MAX_CHUNKS:
            raise RagError(f"More than {MAX_CHUNKS:,} passages. Split the document first.")
    for i, c in enumerate(out):
        c["idx"] = i
    return out


def _split_long(para: str) -> list[str]:
    if len(para) <= CHUNK_CHARS:
        return [para]
    sentences = re.split(r"(?<=[.!?])\s+", para)
    out, cur = [], ""
    for s in sentences:
        while len(s) > CHUNK_CHARS:
            out.append(s[:CHUNK_CHARS])
            s = s[CHUNK_CHARS:]
        if cur and len(cur) + len(s) + 1 > CHUNK_CHARS:
            out.append(cur)
            cur = s
        else:
            cur = (cur + " " + s).strip()
    if cur:
        out.append(cur)
    return out


# ------------------------------------------------------------------ embeddings
def embeddings_enabled() -> bool:
    return bool(config.AZURE_OPENAI_EMBEDDING_DEPLOYMENT and config.AZURE_OPENAI_API_KEY and config.AZURE_OPENAI_ENDPOINT)


def _embedder():
    from langchain_openai import OpenAIEmbeddings
    return OpenAIEmbeddings(model=config.AZURE_OPENAI_EMBEDDING_DEPLOYMENT, base_url=config.AZURE_OPENAI_ENDPOINT,
                            api_key=config.AZURE_OPENAI_API_KEY, check_embedding_ctx_length=False, timeout=60)


def embed_texts(texts: list[str]) -> list[list[float]]:
    emb = _embedder()
    out: list[list[float]] = []
    for i in range(0, len(texts), EMBED_BATCH):
        out += emb.embed_documents(texts[i:i + EMBED_BATCH])
    return out


def embed_query(text: str) -> list[float]:
    return _embedder().embed_query(text)


# ------------------------------------------------------------------ storage
def ensure_schema(lk: sqlite3.Connection) -> None:
    lk.execute("""CREATE TABLE IF NOT EXISTS rag_chunks (
        id INTEGER PRIMARY KEY AUTOINCREMENT, dataset_id INTEGER NOT NULL, idx INTEGER NOT NULL,
        page INTEGER, text TEXT NOT NULL, embedding BLOB)""")
    lk.execute("CREATE INDEX IF NOT EXISTS idx_rag_ds ON rag_chunks(dataset_id)")
    lk.execute("CREATE VIRTUAL TABLE IF NOT EXISTS rag_fts USING fts5(text, tokenize='porter unicode61')")


def index(lake_db: Path, dataset_id: int, chunks: list[dict]) -> dict:
    """Store passages (and embeddings when configured). Returns {"semantic": bool, "note": str}."""
    vectors, note = None, ""
    if embeddings_enabled():
        try:
            vectors = embed_texts([c["text"] for c in chunks])
        except Exception as exc:  # noqa: BLE001  (keyword search still works)
            note = f"Embeddings failed, keyword search only: {str(exc)[:160]}"
    with sqlite3.connect(lake_db) as lk:
        ensure_schema(lk)
        for i, c in enumerate(chunks):
            blob = np.asarray(vectors[i], dtype=np.float32).tobytes() if vectors else None
            cur = lk.execute("INSERT INTO rag_chunks (dataset_id, idx, page, text, embedding) VALUES (?,?,?,?,?)",
                             (dataset_id, c["idx"], c["page"], c["text"], blob))
            lk.execute("INSERT INTO rag_fts (rowid, text) VALUES (?, ?)", (cur.lastrowid, c["text"]))
    return {"semantic": vectors is not None, "note": note}


def delete(lake_db: Path, dataset_id: int) -> None:
    if not lake_db.exists():
        return
    with sqlite3.connect(lake_db) as lk:
        ensure_schema(lk)
        ids = [r[0] for r in lk.execute("SELECT id FROM rag_chunks WHERE dataset_id = ?", (dataset_id,))]
        for i in ids:
            lk.execute("DELETE FROM rag_fts WHERE rowid = ?", (i,))
        lk.execute("DELETE FROM rag_chunks WHERE dataset_id = ?", (dataset_id,))


def chunk_rows(lake_db: Path, dataset_id: int, limit: int = 500) -> list[dict]:
    if not lake_db.exists():
        return []
    with sqlite3.connect(lake_db) as lk:
        ensure_schema(lk)
        return [{"passage": r[0], "page": r[1], "text": r[2]} for r in lk.execute(
            "SELECT idx, page, text FROM rag_chunks WHERE dataset_id = ? ORDER BY idx LIMIT ?", (dataset_id, limit))]


def count(lake_db: Path) -> int:
    if not lake_db.exists():
        return 0
    with sqlite3.connect(lake_db) as lk:
        ensure_schema(lk)
        return lk.execute("SELECT COUNT(*) FROM rag_chunks").fetchone()[0]


# ------------------------------------------------------------------ search
def _fts_query(text: str) -> str:
    words = [w for w in re.findall(r"\w{2,}", text.lower())][:16]
    return " OR ".join(f'"{w}"' for w in dict.fromkeys(words))


def search(lake_db: Path, query: str, dataset_ids: list[int], k: int = 6) -> list[dict]:
    """Best passages for the query across the given text datasets: [{dataset_id, page, text, score, mode}]."""
    if not dataset_ids or not lake_db.exists() or not query.strip():
        return []
    marks = ",".join("?" * len(dataset_ids))
    ranked: dict[int, float] = {}
    modes: dict[int, set] = {}
    with sqlite3.connect(lake_db) as lk:
        ensure_schema(lk)
        q = _fts_query(query)
        if q:
            rows = lk.execute(
                f"""SELECT c.id FROM rag_fts f JOIN rag_chunks c ON c.id = f.rowid
                    WHERE rag_fts MATCH ? AND c.dataset_id IN ({marks}) ORDER BY bm25(rag_fts) LIMIT 20""",
                [q, *dataset_ids]).fetchall()
            for rank, (cid,) in enumerate(rows):
                ranked[cid] = ranked.get(cid, 0) + 1 / (60 + rank)
                modes.setdefault(cid, set()).add("keyword")
        if embeddings_enabled():
            try:
                qv = np.asarray(embed_query(query), dtype=np.float32)
                data = lk.execute(f"SELECT id, embedding FROM rag_chunks WHERE embedding IS NOT NULL AND dataset_id IN ({marks})",
                                  dataset_ids).fetchall()
                if data:
                    ids = [d[0] for d in data]
                    mat = np.vstack([np.frombuffer(d[1], dtype=np.float32) for d in data])
                    sims = mat @ qv / (np.linalg.norm(mat, axis=1) * np.linalg.norm(qv) + 1e-9)
                    for rank, j in enumerate(np.argsort(-sims)[:20]):
                        ranked[ids[j]] = ranked.get(ids[j], 0) + 1 / (60 + rank)
                        modes.setdefault(ids[j], set()).add("semantic")
            except Exception:  # noqa: BLE001  (fall back to keyword results)
                pass
        top = sorted(ranked, key=ranked.get, reverse=True)[:k]
        out = []
        for cid in top:
            r = lk.execute("SELECT dataset_id, page, text FROM rag_chunks WHERE id = ?", (cid,)).fetchone()
            out.append({"dataset_id": r[0], "page": r[1], "text": r[2], "score": round(ranked[cid] * 60, 3),
                        "mode": "+".join(sorted(modes[cid]))})
    return out
